from docx import Document
from io import BytesIO

def build_docx(patient_name, patient_dob, visit_date, provider_name, visit_type, soap: dict) -> BytesIO:
    doc = Document()
    doc.add_heading("Clinical SOAP Note", level=1)

    meta = doc.add_paragraph()
    meta.add_run(f"Patient Name: {patient_name or 'N/A'}\n").bold = True
    meta.add_run(f"Date of Birth: {patient_dob}\n")
    meta.add_run(f"Visit Date: {visit_date}\n")
    meta.add_run(f"Provider: {provider_name or 'N/A'}\n")
    meta.add_run(f"Visit Type: {visit_type}\n")

    for section in ["SUBJECTIVE", "OBJECTIVE", "ASSESSMENT", "PLAN"]:
        doc.add_heading(section.title(), level=2)
        doc.add_paragraph(soap.get(section, ""))

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer